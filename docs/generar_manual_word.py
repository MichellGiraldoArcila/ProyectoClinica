"""Genera MANUAL_USUARIO.docx desde MANUAL_USUARIO.md"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

BASE = Path(__file__).parent
md_path = BASE / 'MANUAL_USUARIO.md'
out_path = BASE / 'MANUAL_USUARIO.docx'

with md_path.open(encoding='utf-8') as f:
    lines = f.readlines()

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

t = doc.add_heading('Manual de Usuario', 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
s = doc.add_paragraph('Sistema HCE — IPS Salud y Vida')
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Resolución 866 de 2021 | Versión 1.0 | Mayo 2026')
doc.add_page_break()

in_code = False
for raw in lines:
    line = raw.rstrip()
    if line.startswith('```'):
        in_code = not in_code
        continue
    if in_code:
        doc.add_paragraph(line, style='Intense Quote')
        continue
    if not line or line == '---':
        continue
    if line.startswith('# '):
        continue
    if line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=2)
    elif line.startswith('#### '):
        doc.add_heading(line[5:].strip(), level=3)
    elif line.startswith('|'):
        if '---' in line:
            continue
        doc.add_paragraph(line.replace('|', ' ').strip())
    elif line.startswith('- ') or line.startswith('* '):
        doc.add_paragraph(line[2:].strip(), style='List Bullet')
    elif re.match(r'^\d+\.\s', line):
        doc.add_paragraph(re.sub(r'^\d+\.\s*', '', line), style='List Number')
    else:
        clean = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', line)
        clean = re.sub(r'`([^`]+)`', r'\1', clean)
        clean = clean.replace('**', '')
        if clean.strip():
            doc.add_paragraph(clean.strip())

doc.save(out_path)
print(f'Generado: {out_path}')
